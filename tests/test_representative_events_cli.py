"""The CLI that turns a saved selection into hydrographs, plots and a workbook.

Runs the real script against a synthetic project: a sims list, an mcdf, the
stored hydrographs a URBS run leaves in ``../urbs_results``, and the selection
file the launcher writes. The hyetograph is rebuilt from rainfall data that
lives outside the repo, so it is switched off here and covered separately in
test_event_storm.py; everything else - finding the files, lining the series up
on the main burst, and what gets written - is end to end.
"""

from __future__ import annotations

import importlib.util
import json
import sys
from pathlib import Path

import pandas as pd
import pytest

BRYAN_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BRYAN_ROOT))

pytest.importorskip("scipy")
pytest.importorskip("matplotlib")
pytest.importorskip("openpyxl")

from lib import RepresentativeEvents as events                       # noqa: E402

OUTPUT_FILE = r"sims_mc\results\TFD_mc_24h_GWL1p3"
SIMULATION_PERIOD = 96.0
PREBURST_HOURS = 6.0            # so the run is 102 h long, not 96
SIMS = ("sim_00000", "sim_00001", "sim_00002")


def cli():
    path = BRYAN_ROOT / "util" / "RepresentativeEvents.py"
    spec = importlib.util.spec_from_file_location("RepresentativeEventsCli", path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def write_sims_list(path):
    from openpyxl import Workbook

    columns = ["Include", "Method", "Duration", "Run models", "Analyse results",
               "GWL", "FSL", "Output file", "Output suffix", "Config file",
               "Focal subcatchments", "Exclusions", "IL", "CL"]
    row = ["yes", "monte carlo", 24, "yes", "yes", 1.3, 670.42, OUTPUT_FILE, "",
           "mc_config.json", "focal.csv", "", 30, 2.5]
    workbook = Workbook()
    sheet = workbook.active
    for position, name in enumerate(columns, start=1):
        sheet.cell(row=1, column=position, value=name)
    for position, value in enumerate(row, start=1):
        sheet.cell(row=2, column=position, value=value)
    workbook.save(path)


def hydrograph_frame(scale):
    """A run 102 hours long - the nominal 96 plus a six hour pre-burst."""
    times = [step * 0.5 for step in range(int((SIMULATION_PERIOD + PREBURST_HOURS) / 0.5) + 1)]
    frame = pd.DataFrame({"Time": times})
    for position, name in enumerate(SIMS):
        peak = 30.0 + position * 5
        frame[name] = [scale * (1.0 + hour) / (1.0 + abs(hour - peak)) for hour in times]
    return frame.set_index("Time")


def mcdf_frame():
    rows = {}
    for position in range(len(SIMS)):
        rows[position] = {
            "m": 0, "n": position, "rain_z": 3.09, "rain_aep": 1000.0,
            "mean_rain_mm": 200.0, "tp": 4, "storm_method": "ARR point",
            "preburst_p": 0.5, "preburst_proportion": 0.3, "preburst_mm": 60.0,
            "preburst_tp": 7, "il_p": 0.5, "cl_p": 0.5, "lake_z": 0.0,
            "embedded_bursts": "No embedded bursts", "ADV": 210_000.0,
            "inflow": 3000.0, "level": 220.0 + position, "outflow": 2500.0,
            "level_aep": 0.001, "inflow_aep": 0.001, "outflow_aep": 0.001,
        }
    return pd.DataFrame.from_dict(rows, orient="index")


@pytest.fixture
def project(tmp_path):
    folder = tmp_path / "project"
    results = folder / "sims_mc" / "results"
    stored = folder / "sims_mc" / "urbs_results"
    results.mkdir(parents=True)
    stored.mkdir(parents=True)

    write_sims_list(folder / "sims.xlsx")
    mcdf_frame().to_csv(results / "TFD_mc_24h_GWL1p3__mcdf.csv")
    for kind, scale in (("inflows", 100.0), ("outflows", 80.0), ("levels", 2.0)):
        frame = hydrograph_frame(scale)
        if kind == "levels":
            frame = frame + 218.0
        frame.to_csv(stored / f"TFD_mc_24h_GWL1p3_{kind}.csv")

    (folder / "model_config.json").write_text(json.dumps({
        "model_type": "URBS", "simulation_periods": {"24": SIMULATION_PERIOD}}))
    for name in ("storm_config.json", "climate_config.json"):
        (folder / name).write_text("{}")
    (folder / "sims_config.json").write_text(json.dumps({
        "simulation_list": "sims.xlsx",
        "filepaths": {"model_config": "model_config.json",
                      "storm_config": "storm_config.json",
                      "climate_config": "climate_config.json"},
    }))
    return folder


def write_selection(folder, targets):
    path = folder / "sims_mc" / "results" / "GWL1p3_representative_events.json"
    path.write_text(json.dumps(events.selection_payload(targets, {"group": "GWL1p3"})))
    return path


def target(**overrides):
    settings = dict(kind="aep", value=1000, result_type="level", source="24h",
                    output_file="TFD_mc_24h_GWL1p3",
                    database=r"sims_mc\results\TFD_mc_24h_GWL1p3__mcdf.csv",
                    picked=1)
    settings.update(overrides)
    return events.Target(**settings)


def run(project, selection, *extra):
    module = cli()
    code = module.main(["--config", str(project / "sims_config.json"),
                        "--selection", str(selection), "--no-hyetograph", *extra])
    return module, code


# -- what it writes ----------------------------------------------------------

def test_it_writes_a_workbook_and_a_plot_per_event(project):
    selection = write_selection(project, [target(), target(value=100, picked=2)])
    _, code = run(project, selection)
    assert code == 0

    folder = selection.parent
    workbook = folder / "GWL1p3_representative_events_events.xlsx"
    assert workbook.is_file()
    plots = sorted(path.name for path in folder.glob("*.png"))
    assert plots == ["GWL1p3_representative_events_1_in_100_sim00002.png",
                     "GWL1p3_representative_events_1_in_1_000_sim00001.png"]


def test_the_workbook_holds_a_sheet_per_series(project):
    selection = write_selection(project, [target()])
    run(project, selection)
    workbook = selection.parent / "GWL1p3_representative_events_events.xlsx"
    sheets = pd.read_excel(workbook, sheet_name=None)
    assert {"events", "inflows", "levels", "outflows"} <= set(sheets)
    assert not sheets["inflows"].empty


def test_the_summary_names_the_event_and_its_metrics(project):
    selection = write_selection(project, [target()])
    run(project, selection)
    workbook = selection.parent / "GWL1p3_representative_events_events.xlsx"
    summary = pd.read_excel(workbook, sheet_name="events")
    assert summary.loc[0, "sim"] == 1
    assert summary.loc[0, "hydrograph"] == "sim_00001"
    assert summary.loc[0, "loading"] == "1 in 1,000"
    assert summary.loc[0, "level"] == 221.0          # from the mcdf
    assert summary.loc[0, "rain_aep"] == 1000.0


def test_only_the_chosen_simulation_is_extracted(project):
    selection = write_selection(project, [target(picked=2)])
    run(project, selection)
    workbook = selection.parent / "GWL1p3_representative_events_events.xlsx"
    inflows = pd.read_excel(workbook, sheet_name="inflows")
    assert [name for name in inflows.columns if name.startswith("1_in")] == \
        ["1_in_1_000_sim00002"]


# -- lining the series up on the main burst ----------------------------------

def test_the_series_are_shifted_onto_the_start_of_the_main_burst(project):
    """The run is longer than its simulation period by the pre-burst duration.

    Simulator.run_models does `simulation_period += preburst_duration`, so the
    overshoot is the pre-burst, and taking it off puts t = 0 at the start of the
    main burst - which is where the hyetograph's zero is.
    """
    selection = write_selection(project, [target()])
    module, _ = run(project, selection)
    workbook = selection.parent / "GWL1p3_representative_events_events.xlsx"

    inflows = pd.read_excel(workbook, sheet_name="inflows", index_col=0)
    assert inflows.index.min() == pytest.approx(-PREBURST_HOURS)
    assert inflows.index.max() == pytest.approx(SIMULATION_PERIOD)

    summary = pd.read_excel(workbook, sheet_name="events")
    assert summary.loc[0, "time shift (h)"] == pytest.approx(PREBURST_HOURS)


def test_a_missing_simulation_period_leaves_the_times_alone(project):
    (project / "model_config.json").write_text(json.dumps({"model_type": "URBS"}))
    selection = write_selection(project, [target()])
    run(project, selection)
    workbook = selection.parent / "GWL1p3_representative_events_events.xlsx"
    inflows = pd.read_excel(workbook, sheet_name="inflows", index_col=0)
    # Falls back to twice the duration (48 h), so the overshoot is 54 h.
    assert inflows.index.min() == pytest.approx(-54.0)


# -- the plot ----------------------------------------------------------------

def event_with_rain(project, selection):
    """One collected event, with a hyetograph attached by hand.

    The rainfall data a real rebuild needs lives outside the repo, and the
    panel still has to be drawn correctly.
    """
    module = cli()
    from lib import EventStorm

    config = module.load_config(str(project / "sims_config.json"))
    frame = module.read_sims_list(config["sims_list"])
    targets, _ = events.read_selection(selection)
    event = module.collect(targets[0], frame, config, {}, rebuild=False)

    depths = pd.Series([1.0, 2.0, 8.0, 4.0, 2.0, 1.0],
                       index=[-2.0, -1.0, 0.0, 1.0, 2.0, 3.0])
    event["hyetograph"] = EventStorm.Hyetograph(depths, 1.0, 2.0, [])
    return module, event


def test_the_rainfall_axis_is_reversed(project):
    """Zero at the top, the storm hanging down from it - and a real range.

    `set_ylim(top=0)` alone collapses the axis to nothing, which draws the
    whole storm as one block and looks like a plot until you read the numbers.
    """
    selection = write_selection(project, [target()])
    module, event = event_with_rain(project, selection)

    figure = module.figure_for(event)
    try:
        rain = figure.axes[0]
        bottom, top = rain.get_ylim()
        assert bottom > top                      # reversed
        assert top == pytest.approx(0.0)
        assert bottom >= 8.0                     # covers the deepest interval
    finally:
        module.plt.close(figure)


def test_the_three_panels_share_one_time_axis(project):
    selection = write_selection(project, [target()])
    module, event = event_with_rain(project, selection)

    figure = module.figure_for(event)
    try:
        rain, flows, level = figure.axes[:3]
        assert rain.get_xlim() == flows.get_xlim() == level.get_xlim()
        assert "Flow" in flows.get_ylabel()
        assert "Lake level" in level.get_ylabel()
        assert "main burst" in level.get_xlabel()
    finally:
        module.plt.close(figure)


def test_the_full_supply_level_is_drawn_when_the_row_has_one(project):
    selection = write_selection(project, [target()])
    module, event = event_with_rain(project, selection)

    figure = module.figure_for(event)
    try:
        level = figure.axes[2]
        drawn = [line.get_ydata()[0] for line in level.get_lines()
                 if len(set(line.get_ydata())) == 1]
        assert pytest.approx(670.42) in drawn
    finally:
        module.plt.close(figure)


# -- what it does when something is missing ----------------------------------

def test_a_missing_hydrograph_file_is_a_note_not_a_crash(project):
    (project / "sims_mc" / "urbs_results" / "TFD_mc_24h_GWL1p3_levels.csv").unlink()
    selection = write_selection(project, [target()])
    _, code = run(project, selection)
    assert code == 0

    workbook = selection.parent / "GWL1p3_representative_events_events.xlsx"
    sheets = pd.read_excel(workbook, sheet_name=None)
    assert "levels" not in sheets
    assert "no levels" in str(pd.read_excel(workbook, sheet_name="events").loc[0, "notes"])


def test_a_selection_with_nothing_chosen_says_what_to_do(project, capsys):
    selection = write_selection(project, [target(picked=None)])
    _, code = run(project, selection)
    assert code == 1
    assert "Events page" in capsys.readouterr().out


def test_a_simulation_that_is_not_in_the_database_is_reported(project):
    selection = write_selection(project, [target(picked=99)])
    _, code = run(project, selection)
    assert code == 0
    workbook = selection.parent / "GWL1p3_representative_events_events.xlsx"
    notes = str(pd.read_excel(workbook, sheet_name="events").loc[0, "notes"])
    assert "not in" in notes


def test_the_plots_can_be_skipped(project):
    selection = write_selection(project, [target()])
    run(project, selection, "--no-plots")
    assert not list(selection.parent.glob("*.png"))
    assert (selection.parent / "GWL1p3_representative_events_events.xlsx").is_file()


# -- where the storm inputs come from ----------------------------------------
#
# A reservoir routing row leaves Duration and Focal subcatchments blank by
# design - it re-routes hydrographs a previous run stored - so the storm behind
# a realisation belongs to the run its Input MCDF names. Rebuilding off the
# routed row instead reached pd.read_csv(None) and reported a NoneType buffer.

def sims_frame(rows):
    return pd.DataFrame(rows)


ROUTED = {"Method": "reservoir routing", "Output file": r"routed\TFD_rr_FSL672",
          "Input MCDF": r"sims_mc\results\TFD_mc_24h_GWL1p3__mcdf.csv",
          "Duration": None, "Focal subcatchments": None, "Output suffix": "_FSL672"}
SOURCE = {"Method": "monte carlo", "Output file": OUTPUT_FILE, "Input MCDF": None,
          "Duration": 24, "Focal subcatchments": "focal.csv", "Output suffix": ""}


def test_a_routed_row_takes_its_storm_inputs_from_the_run_it_re_routed():
    module = cli()
    frame = sims_frame([SOURCE, ROUTED])
    row, note = module.storm_row(frame, frame.loc[1])
    assert row["Output file"] == OUTPUT_FILE
    assert row["Duration"] == 24
    assert "TFD_mc_24h_GWL1p3" in note


def test_a_monte_carlo_row_is_its_own_storm_row():
    module = cli()
    frame = sims_frame([SOURCE, ROUTED])
    row, note = module.storm_row(frame, frame.loc[0])
    assert row["Output file"] == OUTPUT_FILE
    assert note is None


def test_a_routed_row_whose_source_is_not_in_the_sims_list_says_so():
    module = cli()
    frame = sims_frame([ROUTED])
    row, note = module.storm_row(frame, frame.loc[0])
    assert row is None
    assert "TFD_mc_24h_GWL1p3__mcdf.csv" in note
    assert note.count("\\") == 0, "the note should name the file, not the path"


def test_a_row_with_no_focal_subcatchments_names_the_key(project):
    module = cli()
    config = module.load_config(str(project / "sims_config.json"))
    row = pd.Series({"Output file": OUTPUT_FILE, "Duration": 24,
                     "Focal subcatchments": None})
    with pytest.raises(ValueError) as raised:
        module.rebuild_hyetograph(row, None, config, {})
    assert "Focal subcatchments" in str(raised.value)


def test_the_focal_file_has_to_exist(project):
    module = cli()
    config = module.load_config(str(project / "sims_config.json"))
    row = pd.Series({"Output file": OUTPUT_FILE, "Duration": 24,
                     "Focal subcatchments": "no_such_focal.csv"})
    with pytest.raises(ValueError) as raised:
        module.rebuild_hyetograph(row, None, config, {})
    assert "not found" in str(raised.value)


def test_a_routed_row_can_name_the_storm_inputs_itself():
    """The realisation is all there in the inherited database.

    A routed mcdf is the inherited one with the routed peaks written over it,
    so every draw the storm was made from survives. Only Duration and Focal
    subcatchments are missing, and a routing row is free to carry them.
    """
    module = cli()
    row = dict(ROUTED, Duration=24, **{"Focal subcatchments": "focal.csv"})
    frame = sims_frame([row])
    found, note = module.storm_row(frame, frame.loc[0])
    assert found is not None
    assert found["Duration"] == 24
    assert "inherited" in note


def test_a_routed_row_with_nothing_to_go_on_says_what_would_fix_it():
    module = cli()
    frame = sims_frame([ROUTED])
    found, note = module.storm_row(frame, frame.loc[0])
    assert found is None
    assert "Focal subcatchments" in note and "--source-sims-list" in note


def test_the_source_run_can_be_in_another_sims_list(project, capsys):
    """Routing rows commonly live in a sims list of their own."""
    other = project / "sources.xlsx"
    write_sims_list(other)                        # holds the monte carlo row

    selection = write_selection(project, [target()])
    module = cli()
    code = module.main(["--config", str(project / "sims_config.json"),
                        "--selection", str(selection), "--no-hyetograph",
                        "--source-sims-list", str(other)])
    assert code == 0
    assert "also looking in" in capsys.readouterr().out
